
import streamlit as st
import fitz  # PyMuPDF
import re
from fpdf import FPDF
from datetime import datetime
from io import BytesIO
import os
from PyPDF2 import PdfReader, PdfWriter
from docx import Document
import tempfile
import shutil
import subprocess
import pandas as pd

po_summary_data = []

def generate_invoice_manual(data: dict, invoice_number: str) -> BytesIO:
    """
    Build an invoice PDF without a PO, then merge with a smart lien waiver PDF.
    Expected data keys:
      customer_name, customer_addr1, customer_addr2,
      job, lot, description, amount (string/number),
      job_location (for waiver), terms (e.g., NET30),
      through_date ('MM/DD/YYYY'), signature (e.g., 'LM')
    """
    # Normalize
    amount_str = str(data.get("amount", "0")).replace(",", "").replace("$", "").strip() or "0"
    try:
        amount_val = float(amount_str)
    except:
        amount_val = 0.0

    invoice_pdf = FPDF()
    invoice_pdf.add_page()
    invoice_pdf.set_font("Arial", size=12)

    # Header
    invoice_pdf.set_font("Arial", 'B', 16)
    invoice_pdf.cell(200, 10, txt="INVOICE", ln=1, align="C")
    invoice_pdf.set_font("Arial", size=12)
    invoice_pdf.cell(100, 8, txt=f"Invoice #: {invoice_number}", ln=0)
    invoice_pdf.cell(100, 8, txt=f"Invoice Date: {datetime.today().strftime('%m/%d/%Y')}", ln=1)
    invoice_pdf.cell(100, 8, txt=f"Terms: {data.get('terms','NET30')}", ln=1)
    invoice_pdf.ln(2)

    # From (your business)
    invoice_pdf.set_font("Arial", 'B', 14)
    invoice_pdf.cell(200, 8, txt="I'll Klean It", ln=1)
    invoice_pdf.set_font("Arial", size=11)
    invoice_pdf.cell(200, 6, txt="Payable to: Jose Moreno", ln=1)
    invoice_pdf.ln(3)

    # Customer block
    invoice_pdf.set_font("Arial", 'B', 12)
    invoice_pdf.cell(200, 7, txt="Bill To:", ln=1)
    invoice_pdf.set_font("Arial", size=11)
    invoice_pdf.cell(200, 6, txt=data.get("customer_name",""), ln=1)
    if data.get("customer_addr1"): invoice_pdf.cell(200, 6, txt=data["customer_addr1"], ln=1)
    if data.get("customer_addr2"): invoice_pdf.cell(200, 6, txt=data["customer_addr2"], ln=1)
    invoice_pdf.ln(3)

    # Job/Lot (optional display for context)
    meta_line = " | ".join([p for p in [
        f"Project: {data.get('job')}" if data.get("job") else None,
        f"Lot: {data.get('lot')}" if data.get("lot") else None
    ] if p])
    if meta_line:
        invoice_pdf.set_font("Arial", 'I', 11)
        invoice_pdf.cell(200, 6, txt=meta_line, ln=1)
        invoice_pdf.ln(2)

    # Table
    invoice_pdf.set_font("Arial", 'B', 12)
    invoice_pdf.cell(120, 8, txt="Description", border=1)
    invoice_pdf.cell(40, 8, txt="Amount", border=1, ln=1)
    invoice_pdf.set_font("Arial", size=11)
    desc = data.get("description", "Services rendered outside scope")
    # split long description into rows
    for line in (desc.split("\n") or [""]):
        invoice_pdf.cell(120, 8, txt=line[:80], border=1)
        # amount only on first row
        invoice_pdf.cell(40, 8, txt=f"${amount_val:,.2f}" if line == (desc.split("\n")[0]) else "", border=1, ln=1)

    # Total
    invoice_pdf.cell(120, 8, txt="", border=0)
    invoice_pdf.set_font("Arial", 'B', 12)
    invoice_pdf.cell(40, 8, txt=f"${amount_val:,.2f}", border=1, ln=1)

    invoice_pdf.ln(8)
    invoice_pdf.set_font("Arial", 'B', 12)
    invoice_pdf.cell(200, 8, txt="THANK YOU FOR YOUR BUSINESS!", ln=1, align="C")

    # Signature line (optional visual)
    invoice_pdf.set_font("Courier", 'I', 16)
    invoice_pdf.cell(200, 12, txt=data.get("signature",""), ln=1, align="C")

    # Bytes
    output_str = invoice_pdf.output(dest='S').encode('latin1')
    inv_buf = BytesIO(output_str)

    # Waiver
    waiver_pdf = generate_waiver_pdf_smart(
        data.get("job_location", "Unknown"),
        f"{amount_val:,.2f}",
        data.get("through_date") or datetime.today().strftime('%m/%d/%Y'),
        signature=data.get("signature","LM")
    )

    # Merge
    result_pdf = fitz.open()
    result_pdf.insert_pdf(fitz.open(stream=inv_buf.getvalue(), filetype="pdf"))
    result_pdf.insert_pdf(fitz.open(stream=waiver_pdf.getvalue(), filetype="pdf"))
    final_buffer = BytesIO()
    result_pdf.save(final_buffer)
    return final_buffer


def generate_waiver_pdf_smart(job_location, amount, through_date, signature="LM"):
    from docx import Document

    template_path = "waiver_template.docx"
    doc = Document(template_path)

    replacements = {
        "{{job_location}}": job_location,
        "{{through_date}}": through_date,
        "{{amount}}": f"${amount}",
        "{{signature}}": signature,
        "{{signature_date}}": through_date,
    }

    for para in doc.paragraphs:
        for key, val in replacements.items():
            if key in para.text:
                for run in para.runs:
                    if key in run.text:
                        run.text = run.text.replace(key, val)

    # Save filled docx
    temp_dir = tempfile.mkdtemp()
    filled_docx_path = os.path.join(temp_dir, "waiver_filled.docx")
    filled_pdf_path = os.path.join(temp_dir, "waiver_filled.pdf")
    doc.save(filled_docx_path)

    # Convert using LibreOffice CLI
    subprocess.run([
        "libreoffice", "--headless", "--convert-to", "pdf", filled_docx_path,
        "--outdir", temp_dir
    ], check=True)

    with open(filled_pdf_path, "rb") as f:
        waiver_pdf = BytesIO(f.read())

    shutil.rmtree(temp_dir)
    return waiver_pdf


# ─── Helper to Extract PO Info ─────────────────────────────────────────────
def extract_po_details(pdf_file):
    doc = fitz.open(stream=pdf_file.read(), filetype="pdf")
    text = "\n".join([page.get_text() for page in doc])
    # Try to extract job location address from top right
    location_match = re.search(r"Lot:\s*\d+\s*\n(.*?)\n(Fresno|Clovis), CA \d{5}", text)
    if location_match:
        job_location = f"{location_match.group(1).strip()}\n{location_match.group(2)}, CA"
    else:
        job_location = "Unknown"

    po_number = None
    match = re.search(r"[A-Za-z0-9]{4,}\s*-\s*[A-Za-z0-9]{1,}\s*-\s*\d{6}", text, re.IGNORECASE)
    if match:
        po_number = match.group(0).replace(" ", "")

    if not po_number:
        for page in doc:
            lines = page.get_text().splitlines()
            for i, line in enumerate(lines):
                if "Purchase Order" in line:
                    after = line.split("Purchase Order")[-1]
                    match = re.search(r"[A-Za-z0-9]{4,}\s*-\s*[A-Za-z0-9]{1,}\s*-\s*\d{6}", after, re.IGNORECASE)
                    if match:
                        po_number = match.group(0).replace(" ", "")
                        break
                    for offset in range(1, 4):
                        if i + offset < len(lines):
                            next_line = lines[i + offset].strip()
                            match = re.search(r"[A-Za-z0-9]{4,}\s*-\s*[A-Za-z0-9]{1,}\s*-\s*\d{6}", next_line, re.IGNORECASE)
                            if match:
                                po_number = match.group(0).replace(" ", "")
                                break
                    if po_number:
                        break
            if po_number:
                break

    details = {
        "po_number": po_number,
        "job_info": re.search(r"Project:\s*(.*?)\nLot:\s*(.*?)\n", text),
        "description": re.search(r"Craft:\s*4440\s*-\s*(.*?)\n", text),
        "amount": re.search(r"Total:\s*\$?([0-9,.]+)", text),
        "customer": re.search(r"Granville Homes Inc\.", text),
    }

    result = {}
    if details["po_number"]:
        result["po_number"] = details["po_number"]
    if details["job_info"]:
        result["job"] = details["job_info"].group(1).strip()
        result["lot"] = details["job_info"].group(2).strip()
    if details["description"]:
        result["description"] = details["description"].group(1).strip()
    if details["amount"]:
        result["amount"] = details["amount"].group(1).strip()
    if details["customer"]:
        result["customer"] = "Granville Homes"
    
    result["job_location"] = job_location

    return result, doc

# ─── Invoice Number Generator ──────────────────────────────────────────────
def get_next_invoice_number():
    counter_file = "invoice_counter.txt"
    if not os.path.exists(counter_file):
        with open(counter_file, "w") as f:
            f.write("1001")
    with open(counter_file, "r+") as f:
        current = int(f.read().strip())
        next_number = current + 1
        f.seek(0)
        f.write(str(next_number))
        f.truncate()
    return f"INV-{current}"

# ─── PDF Invoice Generator ─────────────────────────────────────────────────
def generate_invoice(data, original_po, invoice_number):
    through_date = datetime.today().strftime('%m/%d/%Y')

    waiver_pdf = generate_waiver_pdf_smart(
        data.get("job_location", "Unknown"),
        data.get("amount", "0.00"),
        through_date
    )

    invoice_pdf = FPDF()
    invoice_pdf.add_page()
    invoice_pdf.set_font("Arial", size=12)

    invoice_pdf.set_font("Arial", 'B', 16)
    invoice_pdf.cell(200, 10, txt="INVOICE", ln=1, align="C")
    invoice_pdf.set_font("Arial", size=12)
    invoice_pdf.cell(100, 10, txt=f"Invoice #: {invoice_number}", ln=0)
    invoice_pdf.cell(100, 10, txt=f"Invoice Date: {datetime.today().strftime('%m/%d/%Y')}", ln=1)
    invoice_pdf.cell(100, 10, txt=f"Terms: NET30", ln=1)

    invoice_pdf.set_font("Arial", 'B', 14)
    invoice_pdf.cell(200, 10, txt="I'll Klean It", ln=1, align="L")

    invoice_pdf.set_font("Arial", size=12)
    invoice_pdf.cell(200, 10, txt="Customer: Granville Homes", ln=1)
    invoice_pdf.cell(200, 10, txt="1396 W Herndon", ln=1)
    invoice_pdf.cell(200, 10, txt="Fresno, CA 93711", ln=1)
    invoice_pdf.ln(5)

    invoice_pdf.set_font("Arial", 'B', 12)
    invoice_pdf.cell(60, 10, txt="PO#", border=1)
    invoice_pdf.cell(80, 10, txt="Description", border=1)
    invoice_pdf.cell(40, 10, txt="Amount", border=1, ln=1)

    invoice_pdf.set_font("Arial", size=12)
    invoice_pdf.cell(60, 10, txt=data.get('po_number', 'N/A'), border=1)
    invoice_pdf.cell(80, 10, txt=data.get('description', 'Interior Cleaning'), border=1)
    invoice_pdf.cell(40, 10, txt=f"${data.get('amount', '0.00')}", border=1, ln=1)

    invoice_pdf.cell(140, 10, txt="", border=0)
    invoice_pdf.cell(40, 10, txt=f"${data.get('amount', '0.00')}", border=1, ln=1)

    invoice_pdf.ln(10)
    invoice_pdf.set_font("Arial", 'B', 12)
    invoice_pdf.cell(200, 10, txt="THANK YOU FOR YOUR BUSINESS!", ln=1, align="C")

    invoice_pdf.set_font("Courier", 'I', 18)
    invoice_pdf.cell(200, 20, txt="Luis Moreno", ln=1, align="C")

    output_str = invoice_pdf.output(dest='S').encode('latin1')
    buffer = BytesIO(output_str)
    
    result_pdf = fitz.open()
    result_pdf.insert_pdf(fitz.open(stream=buffer, filetype="pdf"))  # Invoice
    result_pdf.insert_pdf(original_po)                               # PO
    result_pdf.insert_pdf(fitz.open(stream=waiver_pdf.getvalue(), filetype="pdf"))  # Waiver

    final_buffer = BytesIO()
    result_pdf.save(final_buffer)
    return final_buffer

# ─── Streamlit UI ──────────────────────────────────────────────────────────
st.title("📄 Invoice Generator & PO Summary")


tab1, tab2, tab3 = st.tabs(["🧾 Generate Invoices", "📊 PO Summary CSV", "📝 Manual Invoice (No PO)"])


with tab1:
    uploaded_files = st.file_uploader("Upload Client PO(s) (PDF)", type="pdf", accept_multiple_files=True)
    combined_invoice = fitz.open()

    if uploaded_files:
        for uploaded_file in uploaded_files:
            with st.spinner(f"Processing {uploaded_file.name}..."):
                extracted, original_po = extract_po_details(uploaded_file)

            st.subheader(f"Extracted Info from {uploaded_file.name}")
            st.json(extracted)

            st.subheader("Manual Edits (Optional)")
            po_number = st.text_input(f"PO Number for {uploaded_file.name}", value=extracted.get("po_number", ""), key=f"po_{uploaded_file.name}")
            description = st.text_input(f"Description for {uploaded_file.name}", value=extracted.get("description", ""), key=f"desc_{uploaded_file.name}")
            amount = st.text_input(f"Amount ($) for {uploaded_file.name}", value=extracted.get("amount", ""), key=f"amt_{uploaded_file.name}")

            manual_data = {
                "po_number": po_number,
                "description": description,
                "amount": amount,
                "job_location": extracted.get("job_location", "Unknown")
            }
            invoice_number = get_next_invoice_number()
            pdf = generate_invoice(manual_data, original_po, invoice_number)
            combined_invoice.insert_pdf(fitz.open(stream=pdf.getvalue(), filetype="pdf"))

            # Log for summary tab
            po_summary_data.append({
                "Invoice Number": invoice_number,
                "PO Number": po_number,
                "Job": extracted.get("job", ""),
                "Lot": extracted.get("lot", ""),
                "Description": description,
                "Amount": float(amount.replace(",", "").replace("$", "") or 0),
            })

        final_batch = BytesIO()
        combined_invoice.save(final_batch)
        st.download_button("Download Combined Invoice Batch PDF", data=final_batch, file_name="Combined_Invoices.pdf")


with tab2:
    st.subheader("📊 PO Summary & Payment Tracking")
    if po_summary_data:
        df_summary = pd.DataFrame(po_summary_data)
        df_summary["Amount"] = df_summary["Amount"].round(2)
        total = df_summary["Amount"].sum()

        st.dataframe(df_summary, use_container_width=True)
        st.markdown(f"**Total Billed: ${total:,.2f}**")

        csv_buffer = BytesIO()
        df_summary.to_csv(csv_buffer, index=False)
        st.download_button("⬇️ Download PO Summary CSV", csv_buffer.getvalue(), file_name="PO_Summary.csv", mime="text/csv")
    else:
        st.info("No PO summary available yet. Process files in the 'Generate Invoices' tab first.")


with tab3:
    st.subheader("Create Invoice + Lien Waiver (No PO)")
    with st.form("manual_invoice_form", clear_on_submit=False):
        # Customer & Address
        c1, c2 = st.columns(2)
        with c1:
            customer_name = st.text_input("Customer Name", value="Granville Homes")
            job = st.text_input("Project / Job (optional)", value="")
            description = st.text_area("Description", value="Services rendered outside scope")
        with c2:
            customer_addr1 = st.text_input("Customer Address Line 1", value="1396 W Herndon")
            customer_addr2 = st.text_input("Customer Address Line 2", value="Fresno, CA 93711")
            lot = st.text_input("Lot (optional)", value="")

        # Financials & Waiver
        c3, c4, c5 = st.columns(3)
        with c3:
            amount = st.text_input("Amount ($)", value="0.00")
        with c4:
            terms = st.text_input("Terms", value="NET30")
        with c5:
            through_date = st.text_input("Waiver Through Date (MM/DD/YYYY)", value=datetime.today().strftime('%m/%d/%Y'))

        job_location = st.text_input("Job Location (for waiver body)", value=(job or "Unknown"))
        signature = st.text_input("Signature (printed on invoice & waiver)", value="Luis Moreno")

        add_to_summary = st.checkbox("Also add this to the PO Summary tab", value=True)

        submitted = st.form_submit_button("Generate Invoice + Waiver PDF")
        if submitted:
            manual_data = {
                "customer_name": customer_name.strip(),
                "customer_addr1": customer_addr1.strip(),
                "customer_addr2": customer_addr2.strip(),
                "job": job.strip(),
                "lot": lot.strip(),
                "description": description.strip(),
                "amount": amount.strip(),
                "terms": terms.strip(),
                "through_date": through_date.strip(),
                "job_location": job_location.strip(),
                "signature": signature.strip() or "LM",
            }
            invoice_number = get_next_invoice_number()
            pdf_buf = generate_invoice_manual(manual_data, invoice_number)

            st.download_button(
                "⬇️ Download Invoice+Waiver PDF",
                data=pdf_buf.getvalue(),
                file_name=f"{invoice_number}_manual.pdf",
                mime="application/pdf"
            )

            # Optional: log into summary
            if add_to_summary:
                try:
                    amt_f = float(str(amount).replace(",", "").replace("$", "") or 0)
                except:
                    amt_f = 0.0
                po_summary_data.append({
                    "Invoice Number": invoice_number,
                    "PO Number": "",  # none
                    "Job": job.strip(),
                    "Lot": lot.strip(),
                    "Description": description.strip(),
                    "Amount": round(amt_f, 2),
                })
                st.success("Added to PO Summary table.")

